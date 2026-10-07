#!/usr/bin/env python3
"""sync_preprints.py — 논문 추적표(paper_submission_tracker.xlsx)와 홈페이지 preprints.json을 대조·동기화한다.

  check   차이와 미등록 항목을 JSON으로 출력(쓰기 없음). 훅이 사용.
  apply   추적표 상태를 preprints.json에 반영(status/journal 필드만). 신규 항목 추가·게재 전환은 하지 않고 보고만 한다.
  hook    Claude Code SessionStart/Stop 훅. Dropbox 세션에서 차이가 있으면 1회 알린다.
  mark-synced  동기화·배포 완료를 기록한다(현재 지문).

상태 매핑(추적표 Status → 사이트): In preparation→In preparation / Submitted·Under Review·Revision requested·Revised submitted→Under review /
Accepted→Accepted / Rejected→In preparation(재투고 준비) / Published·Withdrawn→보고만(journals.json 이동은 website-updater가 처리).
"""
import argparse, hashlib, json, os, sys, unicodedata
from pathlib import Path

HOME = Path.home()
TRACKER = HOME / "Dropbox/2_Areas/240_Journal/paper_submission_tracker.xlsx"
SITE = HOME / "Dropbox/claude/varde80.github.io"
PREPRINTS = SITE / "src/data/preprints.json"
PROJECTS = HOME / "Dropbox/1_Projects"
STATE_DIR = HOME / "Library/Application Support/claude-hub/website-sync"
STATE = STATE_DIR / "state.json"
MAP = {"in preparation": "In preparation", "submitted": "Under review", "under review": "Under review", "revision requested": "Under review",
       "revised submitted": "Under review", "accepted": "Accepted", "rejected": "In preparation"}
REPORT_ONLY = {"published", "withdrawn"}
# Website ID가 n/a 이면 홈페이지 대상이 아니다(학회 발표는 conferences.json 소관).

def nfc(s): return unicodedata.normalize("NFC", str(s or ""))

def read_tracker():
    from openpyxl import load_workbook
    ws = load_workbook(TRACKER, read_only=True, data_only=True)["process"]
    rows = list(ws.iter_rows(min_row=1, values_only=True)); head = [nfc(h) for h in rows[0]]
    out = []
    for r in rows[1:]:
        if not r or r[0] is None: continue
        d = dict(zip(head, r)); out.append({"status": nfc(d.get("Status")).strip(), "title": nfc(d.get("Title")).strip(),
                                            "web": nfc(d.get("Website ID")).strip(), "folder": nfc(d.get("Folder")).strip(), "journal": nfc(d.get("Journal / Venue")).strip()})
    return out

def diff():
    prep = json.loads(PREPRINTS.read_text(encoding="utf-8")); by_id = {p["id"]: p for p in prep}
    changes, report, unregistered = [], [], []
    for row in read_tracker():
        st = row["status"].lower()
        if st in REPORT_ONLY:
            if row["web"].startswith("prep"): report.append({"id": row["web"], "tracker_status": row["status"], "note": "게재/철회 — preprints에서 제거하고 journals.json로 이동 필요"})
            continue
        target = MAP.get(st)
        if not target: report.append({"id": row["web"], "tracker_status": row["status"], "note": "매핑 없는 상태값"}); continue
        if row["web"].lower() in ("n/a", "-", "없음"): continue
        if not row["web"].startswith("prep"):
            unregistered.append({"title": row["title"][:80], "tracker_status": row["status"], "folder": row["folder"]}); continue
        p = by_id.get(row["web"])
        if not p: report.append({"id": row["web"], "note": "preprints.json에 없는 ID"}); continue
        if p.get("status") != target or p.get("journal") != target:
            changes.append({"id": row["web"], "from": p.get("status"), "to": target, "title": p.get("title", "")[:70]})
    tracked_folders = {nfc(r["folder"]) for r in read_tracker()}
    new_folders = []
    if PROJECTS.is_dir():
        for d in PROJECTS.iterdir():
            n = nfc(d.name)
            if d.is_dir() and "_paper_" in n and not any(n in f for f in tracked_folders): new_folders.append(n)
    return {"changes": changes, "report": report, "unregistered": unregistered, "new_paper_folders": sorted(new_folders)}

def fingerprint(result):
    return hashlib.sha256(json.dumps(result, ensure_ascii=False, sort_keys=True).encode()).hexdigest()

def apply():
    result = diff(); prep = json.loads(PREPRINTS.read_text(encoding="utf-8"))
    for c in result["changes"]:
        for p in prep:
            if p["id"] == c["id"]: p["status"] = c["to"]; p["journal"] = c["to"]
    PREPRINTS.write_text(json.dumps(prep, ensure_ascii=False, indent=4) + "\n", encoding="utf-8")
    return result

def load_state():
    try: return json.loads(STATE.read_text())
    except Exception: return {}

def save_state(st):
    STATE_DIR.mkdir(parents=True, exist_ok=True); STATE.write_text(json.dumps(st, ensure_ascii=False, indent=1))

def pending(result):
    return bool(result["changes"] or result["unregistered"] or result["new_paper_folders"] or result["report"])

def hook(payload):
    event = payload.get("hook_event_name")
    if payload.get("agent_id") or event not in {"SessionStart", "Stop"} or not payload.get("cwd"): return None
    try: Path(payload["cwd"]).resolve().relative_to((HOME / "Dropbox").resolve())
    except ValueError: return None
    if not TRACKER.exists() or not PREPRINTS.exists(): return None
    result = diff()
    if not pending(result): return None
    fp = fingerprint(result); st = load_state()
    if st.get("synced_fingerprint") == fp: return None
    msg = ("홈페이지 논문 목록 갱신 대기(사용자 상시 승인 2026-10-07: 추적표·논문 폴더 변화 시 자동 반영). "
           "python3 " + str(SITE / "tools/sync_preprints.py") + " check 로 차이를 확인하고, 상태 변경은 `apply`, 신규 항목(unregistered·new_paper_folders)은 "
           "varde80.github.io의 website-updater 절차로 preprints.json에 추가(저자 마커 규칙 준수)한 뒤 `npm run build` 통과 시 commit·push. "
           "완료 후 `python3 " + str(SITE / "tools/sync_preprints.py") + " mark-synced` 실행. 사용자가 중단한 작업을 다시 시작하지 말 것. 요약: " + json.dumps(result, ensure_ascii=False)[:900])
    if event == "SessionStart":
        return {"hookSpecificOutput": {"hookEventName": event, "additionalContext": msg}}
    key = hashlib.sha256((str(payload.get("session_id", "")) + fp).encode()).hexdigest()
    nudges = st.setdefault("nudges", [])
    if payload.get("stop_hook_active") or key in nudges: return None
    nudges.append(key); st["nudges"] = nudges[-128:]; save_state(st)
    return {"decision": "block", "reason": msg}

def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("command", choices=("check", "apply", "hook", "mark-synced")); a = ap.parse_args(argv)
    try:
        if a.command == "check": print(json.dumps(diff(), ensure_ascii=False, indent=1))
        elif a.command == "apply": print(json.dumps(apply(), ensure_ascii=False, indent=1))
        elif a.command == "mark-synced":
            st = load_state(); st["synced_fingerprint"] = fingerprint(diff()); save_state(st); print(json.dumps({"synced_fingerprint": st["synced_fingerprint"]}))
        else:
            r = hook(json.loads(sys.stdin.read(65536)))
            if r is not None: print(json.dumps(r, ensure_ascii=False))
        return 0
    except Exception as e:
        print("website-sync: 보류 (" + type(e).__name__ + ": " + str(e)[:80] + ")", file=sys.stderr); return 0 if a.command == "hook" else 1

if __name__ == "__main__": sys.exit(main())
