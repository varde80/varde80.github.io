#!/usr/bin/env python3
"""
Editable DOCX twin of generate_cv.py.

Same data files, same filters (--doi-list / --if-override / --grants-from …) and the
same wording — every publication line, project line, label and statistic is produced by
the functions in generate_cv.py, so the Word file says exactly what the PDF says.
Only the rendering backend differs (python-docx instead of reportlab).

Usage (mirrors generate_cv.py):
  venv/bin/python generate_cv_docx.py --doi-list DOIS.txt --if-override IF.json \
      --pub-subtitle "(2022.9-2026.9)" --grants-from 2021.09 --suffix _Ajou --lang both
Output: output/{YYYYMMDD}_CV_HLee{_KR}{suffix}.docx  (or --outdir DIR)
"""
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

import generate_cv as G  # single source of truth for data + wording

# ---------------------------------------------------------------------------
# Fonts / colours (hex without '#', as Word wants them)
# ---------------------------------------------------------------------------
LATIN_FONT = "Arial"        # stands in for Helvetica
EA_FONT = "Malgun Gothic"    # 맑은 고딕 — bundled with Office on Windows and macOS (English family name matches everywhere)

NAVY = "2D3748"
ACCENT = "4A5568"
LIGHT_GRAY = "718096"
LINK_BLUE = "2563EB"
HIGHLIGHT_BG = "F0F4F8"
HEADER_TEXT = "E2E8F0"
WHITE = "FFFFFF"
BLACK = "000000"


def _hex(c):
    return c.lstrip('#').upper()


# ---------------------------------------------------------------------------
# Low-level OOXML helpers
# ---------------------------------------------------------------------------
# CT_PPr child order (ECMA-376 §17.3.1.26); Word can reject out-of-order children.
PPR_SEQ = ('w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore', 'w:framePr', 'w:widowControl',
           'w:numPr', 'w:suppressLineNumbers', 'w:pBdr', 'w:shd', 'w:tabs', 'w:suppressAutoHyphens',
           'w:kinsoku', 'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE', 'w:autoSpaceDN',
           'w:bidi', 'w:adjustRightInd', 'w:snapToGrid', 'w:spacing', 'w:ind', 'w:contextualSpacing',
           'w:mirrorIndents', 'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
           'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle', 'w:rPr', 'w:sectPr', 'w:pPrChange')


def _insert_ppr_child(pPr, child):
    """Insert a pPr child at its schema position (Word rejects out-of-order children)."""
    seq = PPR_SEQ
    tag = child.tag.split('}')[1]
    try:
        successors = seq[seq.index('w:' + tag) + 1:]
    except ValueError:
        pPr.append(child)
        return
    pPr.insert_element_before(child, *successors)


def style_run(run, size=None, bold=None, italic=None, color=None, superscript=False,
              latin=LATIN_FONT, ea=EA_FONT):
    rPr = run._r.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(attr), latin)
    rFonts.set(qn('w:eastAsia'), ea)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(_hex(color))
    if superscript:
        run.font.superscript = True


def shade_paragraph(p, fill):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), _hex(fill))
    _insert_ppr_child(pPr, shd)


def bottom_border(p, color, sz=16, space=1):
    """sz in eighths of a point (16 = 2 pt)."""
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement('w:pBdr')
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single')
    b.set(qn('w:sz'), str(sz))
    b.set(qn('w:space'), str(space))
    b.set(qn('w:color'), _hex(color))
    pBdr.append(b)
    _insert_ppr_child(pPr, pBdr)


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), _hex(fill))
    tcPr.append(shd)


def set_table_cell_margins(tbl, top, bottom, left, right):
    """Margins in twips (1 pt = 20)."""
    tblPr = tbl._tbl.tblPr
    mar = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement('w:' + k)
        e.set(qn('w:w'), str(int(v)))
        e.set(qn('w:type'), 'dxa')
        mar.append(e)
    look = tblPr.find(qn('w:tblLook'))
    if look is not None:
        look.addprevious(mar)
    else:
        tblPr.append(mar)


def set_col_widths(tbl, widths):
    """Fixed layout; write widths to both tblGrid/gridCol (LibreOffice, Word) and every tcW."""
    tbl.autofit = False
    for col, w in zip(tbl.columns, widths):
        col.width = w
    for row in tbl.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def wrap_in_hyperlink(paragraph, run, url):
    r_id = paragraph.part.relate_to(url, RT.HYPERLINK, is_external=True)
    h = OxmlElement('w:hyperlink')
    h.set(qn('r:id'), r_id)
    r = run._r
    r.addprevious(h)
    h.append(r)


# ---------------------------------------------------------------------------
# reportlab-style inline markup -> DOCX runs
#   supports <b> <i> <super> <font color= face= size=> <link href=> <br/> and
#   the entities &nbsp; &amp; &lt; &gt; &quot; &#NNN;   (a bare '&' is literal)
# ---------------------------------------------------------------------------
TOKEN_RE = re.compile(r'(<[^<>]+>|&(?:nbsp|amp|lt|gt|quot|#\d+);)')
ATTR_RE = re.compile(r'(\w+)\s*=\s*"([^"]*)"')
ENTITIES = {'&nbsp;': ' ', '&amp;': '&', '&lt;': '<', '&gt;': '>', '&quot;': '"'}


def add_markup(paragraph, markup, size=9, color=None, bold=False, italic=False):
    bold_on, italic_on, super_on, link = bold, italic, False, None
    colors, sizes = [color], [size]
    for tok in TOKEN_RE.split(markup or ''):
        if not tok:
            continue
        if tok.startswith('<'):
            m = re.match(r'</?\s*([a-zA-Z]+)', tok)
            name = m.group(1).lower() if m else ''
            closing = tok.startswith('</')
            if name == 'b':
                bold_on = not closing
            elif name == 'i':
                italic_on = not closing
            elif name == 'super':
                super_on = not closing
            elif name == 'font':
                if closing:
                    if len(colors) > 1:
                        colors.pop()
                        sizes.pop()
                else:
                    attrs = dict(ATTR_RE.findall(tok))
                    colors.append(attrs.get('color', colors[-1]))
                    sizes.append(float(attrs['size']) if attrs.get('size') else sizes[-1])
            elif name == 'link':
                link = None if closing else dict(ATTR_RE.findall(tok)).get('href')
            elif name == 'br':
                paragraph.add_run().add_break(WD_BREAK.LINE)
            continue
        text = ENTITIES.get(tok)
        if text is None and tok.startswith('&#'):
            text = chr(int(tok[2:-1]))
        if text is None:
            text = tok
        for i, piece in enumerate(text.split('\n')):
            if i:
                paragraph.add_run().add_break(WD_BREAK.LINE)
            if piece == '':
                continue
            run = paragraph.add_run(piece)
            style_run(run, size=sizes[-1], bold=bold_on, italic=italic_on,
                      color=colors[-1], superscript=super_on)
            if link:
                wrap_in_hyperlink(paragraph, run, link)


def para(container, markup='', size=9, color=None, bold=False, italic=False, align=None,
         before=0, after=0, leading=None, left=0, first=0, shade=None, keep_next=False,
         reuse_first=False):
    """Add a paragraph (or reuse a fresh cell's empty first paragraph) and fill it."""
    if reuse_first and container.paragraphs and not container.paragraphs[0].text \
            and not container.paragraphs[0].runs:
        p = container.paragraphs[0]
    else:
        p = container.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if leading:
        pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        pf.line_spacing = Pt(leading)
    if left or first:
        pf.left_indent = Pt(left)
        pf.first_line_indent = Pt(first)
    if align is not None:
        pf.alignment = align
    if keep_next:
        pf.keep_with_next = True
    if shade:
        shade_paragraph(p, shade)
    if markup:
        add_markup(p, markup, size=size, color=color, bold=bold, italic=italic)
    return p


def page_break(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_break(WD_BREAK.PAGE)


# ---------------------------------------------------------------------------
# Section renderers (mirror generate_cv.build_* one for one)
# ---------------------------------------------------------------------------
def section_header(doc, ctx, title, icon, subtitle=None, before=16, after=6):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.keep_with_next = True
    text = title if ctx.ko else title.upper()
    r = p.add_run('●')
    style_run(r, size=13, color=NAVY)
    r = p.add_run('  ' + text)
    style_run(r, size=13, bold=True, color=NAVY)
    if subtitle:
        r = p.add_run('  ' + subtitle)
        style_run(r, size=11, color=LIGHT_GRAY)
    return p


def build_title(doc, ctx):
    para(doc, ctx.labels['cv_title'], size=22, bold=True, color=NAVY,
         align=WD_ALIGN_PARAGRAPH.CENTER, after=18)


def header_affiliation(ctx):
    """Same selection rule as generate_cv.create_header_table."""
    if ctx.ko:
        return ctx.kod.get('header_title', ''), ctx.kod.get('header_org', '')
    for exp in ctx.professor.get('experience', []):
        if not isinstance(exp, dict):
            continue
        position, institution, period = exp.get('position', ''), exp.get('institution', ''), exp.get('period', '')
        if ('Director' in position or 'Head' in position) and 'present' in period.lower():
            dept = exp.get('Department', '')
            line2 = f"{dept}<br/>{institution}" if dept else institution.replace(', ', '<br/>')
            return position, line2
    return '', ''


def build_header(doc, ctx):
    L = ctx.labels
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_col_widths(tbl, [Cm(13), Cm(4)])
    set_table_cell_margins(tbl, top=300, bottom=300, left=300, right=300)   # 15 pt
    left, right = tbl.rows[0].cells
    for c in (left, right):
        shade_cell(c, NAVY)
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

    name = ctx.kod['name'] if ctx.ko else ctx.professor['name']
    para(left, name, size=18, color=WHITE, after=4, leading=22, reuse_first=True)
    line1, line2 = header_affiliation(ctx)
    if line1:
        para(left, f"{line1}<br/>{line2}", size=9, color=HEADER_TEXT, after=6, leading=12)
    para(left, f"<b>{L['phone']}</b>  {ctx.professor['phone']}", size=10, color=WHITE, after=1)
    para(left, f"<b>{L['email']}</b>  {ctx.professor['email']}", size=10, color=WHITE, after=1)

    image_path = G.SCRIPT_DIR.parent / "public" / ctx.professor.get("image", "").lstrip("/")
    p = right.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    if image_path.exists():
        p.add_run().add_picture(str(image_path), width=Cm(2.5), height=Cm(3.2))
    else:
        print(f"warning: profile image not found: {image_path}", file=sys.stderr)

    ids = ctx.professor.get('ids') or []
    id_text = "   ·   ".join(f"<b>{d['label']}</b> {d['value']}" for d in ids if d.get('value'))
    p = para(doc, id_text, size=8, color=LIGHT_GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT, before=4, after=2)
    bottom_border(p, NAVY, sz=16, space=4)


def build_summary(doc, ctx):
    bio = ctx.professor.get('bio', '')
    if bio:
        para(doc, bio, size=10, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=10, after=10, leading=14)


def build_interests(doc, ctx):
    if not ctx.professor.get('Research Interests'):
        return
    section_header(doc, ctx, ctx.labels['research_interests'], '◈')
    interests = ctx.kod['Research Interests'] if ctx.ko else ctx.professor['Research Interests']
    for entry in interests:
        if isinstance(entry, dict):
            if entry.get('category'):
                para(doc, entry['category'], size=9.5, bold=True, color=NAVY, before=4, after=1,
                     leading=12, keep_next=True)
            for item in entry.get('items', []):
                para(doc, f"–&nbsp;&nbsp;{item}", size=9, before=0.5, after=0.5, leading=12, left=22, first=-8)
        else:
            para(doc, f"•&nbsp;&nbsp;{entry}", size=9, before=1, after=1, leading=12, left=12, first=-7)


def timeline_table(doc, ctx, rows):
    """rows: [(date_text, title_markup, subtitle_markup|None)] rendered as
    [◆][date][title / subtitle] like generate_cv.create_timeline_entry."""
    tbl = doc.add_table(rows=0, cols=3)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_cell_margins(tbl, top=0, bottom=0, left=0, right=0)
    widths = [Cm(0.6), Cm(2.2), Cm(13)]
    for date_text, title, subtitle in rows:
        cells = tbl.add_row().cells
        for c, w in zip(cells, widths):
            c.width = w
        para(cells[0], '◆', size=7, color=NAVY, before=4, after=4, reuse_first=True)
        para(cells[1], date_text, size=10, color=ACCENT, before=2, after=4, leading=14, reuse_first=True)
        para(cells[2], title, size=10, bold=True, before=2, after=(1 if subtitle else 4), leading=14, reuse_first=True)
        if subtitle:
            para(cells[2], subtitle, size=10, color=ACCENT, italic=not ctx.ko, after=4, leading=13)
    set_col_widths(tbl, widths)
    return tbl


def build_experience(doc, ctx):
    section_header(doc, ctx, ctx.labels['experience'], '◆')

    def exp_start_year(exp):
        period = exp.get('period', '') if isinstance(exp, dict) else (exp.split(', ')[-1] if exp else '')
        return G.first_year(period)

    rows = []
    for exp in sorted(ctx.professor['experience'], key=exp_start_year, reverse=True):
        if isinstance(exp, dict):
            position, department = exp.get('position', ''), exp.get('Department', '')
            institution, period = exp.get('institution', ''), exp.get('period', '')
            if ctx.ko:
                ko_exp = ctx.kod.get('experience', {}).get(period)
                if ko_exp:
                    position, subtitle = ko_exp['position'], ko_exp['org']
                else:
                    subtitle = f"{department}<br/>{institution}" if department else institution
            elif department:
                subtitle = f"{department}<br/>{institution}"
            else:
                subtitle = institution
        else:
            parts = exp.split(', ')
            if len(parts) >= 3:
                position, subtitle, period = parts[0], ', '.join(parts[1:-1]), parts[-1]
            else:
                position, subtitle, period = exp, '', ''
        rows.append((G.format_period_range(period), position, subtitle or None))
    timeline_table(doc, ctx, rows)


def build_education(doc, ctx):
    section_header(doc, ctx, ctx.labels['education'], '◇')
    rows = []
    for edu in ctx.professor['education']:
        if isinstance(edu, dict):
            degree, field = edu.get('degree', ''), edu.get('field', '')
            institution, period = edu.get('institution', ''), edu.get('period', '')
            thesis, advisor = edu.get('thesis', ''), edu.get('advisor', '')
            if ctx.ko:
                ko_edu = ctx.kod.get('education', {}).get(period, {})
                degree = ko_edu.get('degree', degree)
                field = ko_edu.get('field', field)
                institution = ko_edu.get('institution', institution)
                advisor = ko_edu.get('advisor', advisor)
            title = f"{degree}, {field}, {institution}"
            subtitle = None
            if thesis:
                subtitle = f"{thesis} ({advisor})" if advisor else thesis
            elif ctx.ko and advisor:
                subtitle = f"지도교수: {advisor}"
            date_display = G.format_year_range(period)
        else:
            parts = edu.split(', ')
            if len(parts) >= 3:
                title, subtitle, year = parts[0], ', '.join(parts[1:-1]), parts[-1]
            else:
                title, subtitle, year = edu, '', ''
            ys = G.extract_years(year)
            date_display = ys[0] if ys else ''
        rows.append((date_display, title, subtitle or None))
    timeline_table(doc, ctx, rows)


def build_awards(doc, ctx):
    awards_list = ctx.professor.get('Honors and Awards') or ctx.professor.get('Grants and Awards')
    if not awards_list:
        return
    section_header(doc, ctx, ctx.labels['awards'], '★')
    rows = []
    if ctx.ko:
        for a in ctx.kod.get('Honors and Awards', []):
            rows.append((a.get('year', ''), a.get('name', ''), a.get('org') or None))
    else:
        for award in awards_list:
            parts = award.split(', ')
            if len(parts) < 2:
                continue
            date_str, award_parts = '', []
            for part in parts:
                if re.search(r'\d{4}', part):
                    date_str = part
                else:
                    award_parts.append(part)
            award_name = award_parts[0] if award_parts else parts[0]
            institution = ', '.join(award_parts[1:]) if len(award_parts) > 1 else ''
            ys = G.extract_years(date_str)
            rows.append((ys[0] if ys else '', award_name, institution or None))
    timeline_table(doc, ctx, rows)


def build_activities(doc, ctx):
    if not ctx.professor.get('Professional Activities/Memberships'):
        return
    section_header(doc, ctx, ctx.labels['activities'], '●')
    activities = (ctx.kod['Professional Activities/Memberships'] if ctx.ko
                  else ctx.professor['Professional Activities/Memberships'])
    for a in activities:
        para(doc, f"•&nbsp;&nbsp;{a}", size=9, before=2, after=2, leading=12, left=12, first=-7)


def build_publications(doc, ctx):
    L = ctx.labels
    if ctx.section_breaks:
        page_break(doc)
    section_header(doc, ctx, L['publications'], '■', subtitle=(ctx.pub_subtitle or '(2024-2026)'),
                   before=0, after=4)
    para(doc, L['pub_note'], size=8, color=LIGHT_GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT)

    # --- selection & statistics: verbatim logic of generate_cv.build_publications ---
    esci = []
    if ctx.doi_filter:
        key = lambda j: str(j.get('doi', '')).lower().strip()
        recent_journals = [j for j in ctx.journals if key(j) in ctx.doi_filter]
        esci = [j for j in recent_journals if 'esci' in ctx.doi_filter[key(j)]]
    else:
        recent_journals = [j for j in ctx.journals if j.get('year', 0) >= 2024]
    preprint_submitted = [j for j in recent_journals if j.get('status', '').lower() in ['submitted', 'preprint']]
    published_journals = [j for j in recent_journals if not j.get('status')]
    published_journals.sort(key=G.get_pub_sort_key)
    years = sorted(set(p['year'] for p in published_journals), reverse=True)

    journal_corresponding = sum(1 for p in published_journals if G.is_corresponding_author(p['authors']))
    journal_coauthor = len(published_journals) - journal_corresponding
    submitted_corresponding = sum(1 for p in preprint_submitted if G.is_corresponding_author(p['authors']))
    submitted_coauthor = len(preprint_submitted) - submitted_corresponding

    def format_pub_stats(total, corresponding, coauthor):
        parts = [f"{L['total']}: {total}"]
        if corresponding > 0:
            parts.append(f"{L['corresponding']}: {corresponding}")
        if coauthor > 0:
            parts.append(f"{L['coauthor']}: {coauthor}")
        return ", ".join(parts)

    def jcr_pct(pub):
        info = (ctx.if_override or {}).get(str(pub.get('doi', '')).lower().strip()) or ctx.if_data.get(pub.get('journal', '')) or next(
            (v for k, v in ctx.if_data.items() if k.lower() == str(pub.get('journal', '')).lower()), '')
        m = re.search(r'JCR\s*([\d.]+)\s*%', str(info))
        return float(m.group(1)) if m else None

    journal_stats = format_pub_stats(len(published_journals), journal_corresponding, journal_coauthor)
    if esci:
        journal_stats = journal_stats.replace(
            f"{L['total']}: {len(published_journals)}",
            f"{L['total']}: {len(published_journals)} (SCI(E) {len(published_journals) - len(esci)}, ESCI {len(esci)})", 1)
    top10 = [p for p in published_journals if p not in esci and (jcr_pct(p) is not None and jcr_pct(p) <= 10.0)]
    if top10:
        t10_corr = sum(1 for p in top10 if G.is_corresponding_author(p['authors']))
        top10_label = L['top10_sci'] if esci else L['top10']
        journal_stats += f" | {top10_label}: {len(top10)} ({L['corresponding']} {t10_corr}, {L['coauthor']} {len(top10) - t10_corr})"

    para(doc, f"<b>{L['journal_articles']}</b> ({journal_stats})", size=10, bold=False, color=NAVY,
         before=12, after=4, keep_next=True)

    def emit(pub, number, status=None, journal_tag=None):
        rl = G.render_publication(pub, number, ctx, status=status, journal_tag=journal_tag)  # reportlab Paragraph
        highlight = rl.style.name == 'PublicationHighlight'
        para(doc, rl.text, size=9, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=2, after=4, leading=12,
             left=10, first=-10, shade=(HIGHLIGHT_BG if highlight else None))

    n = 1
    for year in years:
        para(doc, f"<b>{year}</b>", size=10, color=LIGHT_GRAY, before=2, after=0, leading=12, keep_next=True)
        for pub in [p for p in published_journals if p['year'] == year]:
            emit(pub, n, journal_tag=('ESCI' if pub in esci else None))
            n += 1

    if preprint_submitted:
        preprint_submitted.sort(key=G.get_pub_sort_key)
        submitted_stats = format_pub_stats(len(preprint_submitted), submitted_corresponding, submitted_coauthor)
        para(doc, f"<b>{L['in_submission']}</b> ({submitted_stats})", size=10, color=NAVY, before=12, after=4, keep_next=True)
        for i, pub in enumerate(preprint_submitted, 1):
            emit(pub, i, status=L['submitted'])


def build_grants(doc, ctx):
    L = ctx.labels
    # --- selection & subtotals: verbatim logic of generate_cv.build_grants ---
    pi_projects = [p for p in ctx.projects if G.is_pi_role(p) and G.get_funding_amount_billion(p) >= 0.1]
    earlier = []
    if ctx.grants_from:
        earlier = sorted([p for p in pi_projects if G.get_project_start_ym(p) < ctx.grants_from],
                         key=G.get_project_start_year, reverse=True)
        pi_projects = [p for p in pi_projects if G.get_project_start_ym(p) >= ctx.grants_from]
    ongoing = sorted([p for p in pi_projects if p.get('status') == 'ongoing'], key=G.get_project_start_year, reverse=True)
    completed = sorted([p for p in pi_projects if p.get('status') == 'completed'], key=G.get_project_start_year, reverse=True)

    total_funding = sum(G.get_funding_amount_billion(p) for p in pi_projects)
    earliest_year = min(G.get_project_start_year(p) for p in pi_projects) if pi_projects else 0
    if total_funding > 0 and ctx.grants_from:
        pi_only = sum(G.get_funding_amount_billion(p) for p in pi_projects
                      if str(p.get('role', {}).get('en', p.get('role', ''))).strip() == 'PI')
        copi = total_funding - pi_only
        yr, mo = ctx.grants_from.split('.')
        if ctx.ko:
            total_funding_str = f"총괄책임 {ctx.fund_amount(pi_only)} · 참여기관 책임 {ctx.fund_amount(copi)}, {yr}.{int(mo)} 이후 수주"
        else:
            total_funding_str = f"PI {ctx.fund_amount(pi_only)} · Co-PI {ctx.fund_amount(copi)}, awarded since {yr}.{int(mo)}"
    elif total_funding > 0:
        total_funding_str = (f"{ctx.fund_amount(total_funding)}, {earliest_year}년 이후" if ctx.ko
                             else f"{ctx.fund_amount(total_funding)} {L['since']} {earliest_year}")
    else:
        total_funding_str = ""

    if ctx.section_breaks:
        page_break(doc)
    section_header(doc, ctx, L['grants'], '◆', subtitle=(f"({total_funding_str})" if total_funding_str else None),
                   before=0, after=4)
    para(doc, L['grant_note'], size=8, color=LIGHT_GRAY, align=WD_ALIGN_PARAGRAPH.RIGHT)

    def group(label, projects):
        para(doc, label, size=10, color=NAVY, before=12, after=4, keep_next=True)
        for proj in projects:
            para(doc, G.format_project_line(proj, ctx), size=9, before=2, after=3, leading=12, left=12, first=-7,
                 shade=(HIGHLIGHT_BG if G.is_large_grant(proj) else None))

    if ongoing:
        t = sum(G.get_funding_amount_billion(p) for p in ongoing)
        group(f"<b>{L['ongoing']}</b> {f'({ctx.fund_amount(t)})' if t > 0 else ''}", ongoing)
    if completed:
        t = sum(G.get_funding_amount_billion(p) for p in completed)
        group(f"<b>{L['completed']}</b> {f'({ctx.fund_amount(t)})' if t > 0 else ''}", completed)
    if earlier:
        yr, mo = ctx.grants_from.split('.')
        t = sum(G.get_funding_amount_billion(p) for p in earlier)
        label = L['earlier_grants'].format(ym=f"{yr}.{int(mo)}")
        group(f"<b>{label}</b> ({ctx.fund_amount(t)})", earlier)


SECTION_BUILDERS = [build_title, build_header, build_summary, build_interests, build_experience,
                    build_education, build_awards, build_activities, build_publications, build_grants]


# ---------------------------------------------------------------------------
# Document setup
# ---------------------------------------------------------------------------
def new_document(ctx):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(1.5)
    sec.top_margin, sec.bottom_margin = Cm(0.5), Cm(1.5)
    normal = doc.styles['Normal']
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(9)
    rPr = normal.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rFonts.set(qn(attr), LATIN_FONT)
    rFonts.set(qn('w:eastAsia'), EA_FONT)
    lang = rPr.find(qn('w:lang'))
    if lang is None:
        lang = OxmlElement('w:lang')
        rPr.append(lang)
    lang.set(qn('w:val'), 'en-US')
    lang.set(qn('w:eastAsia'), 'ko-KR')
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0
    doc.core_properties.title = f"{ctx.labels['cv_title']} — {ctx.professor['name']}"
    doc.core_properties.author = ctx.professor['name']
    return doc


def generate_cv_docx(lang='en', doi_filter=None, pub_subtitle=None, suffix_extra='', grants_from=None,
                     section_breaks=True, if_override=None, outdir=None):
    G.ensure_fonts_registered()     # render_publication builds reportlab Paragraphs (markup validation)
    data = G.load_cv_data()
    G.validate_inputs(data, lang)
    professor = data['professor']
    ctx = G.CVContext(
        lang=lang, professor=professor, kod=professor.get('ko') or {},
        journals=data['journals'], projects=data['projects'], if_data=data['if_data'],
        labels=G.LABELS[lang], styles=G.create_styles(lang), personal=None,
        doi_filter=doi_filter, pub_subtitle=pub_subtitle, grants_from=grants_from,
        section_breaks=section_breaks, if_override=if_override,
    )
    doc = new_document(ctx)
    for build in SECTION_BUILDERS:
        build(doc, ctx)
    out = Path(outdir) if outdir else G.OUTPUT_DIR
    out.mkdir(parents=True, exist_ok=True)
    suffix = ("_KR" if ctx.ko else "") + suffix_extra
    path = out / f"{datetime.now().strftime('%Y%m%d')}_CV_HLee{suffix}.docx"
    doc.save(str(path))
    print(f"DOCX CV generated successfully: {path}")
    return path


def main():
    import argparse
    ap = argparse.ArgumentParser(description="Generate editable DOCX CVs (twin of generate_cv.py).")
    ap.add_argument('--doi-list')
    ap.add_argument('--pub-subtitle')
    ap.add_argument('--suffix', default='')
    ap.add_argument('--lang', choices=['en', 'ko', 'both'], default='both')
    ap.add_argument('--no-section-breaks', action='store_true')
    ap.add_argument('--grants-from')
    ap.add_argument('--if-override')
    ap.add_argument('--outdir', help="output directory (default: cv-generator/output)")
    args = ap.parse_args()

    doi_filter = None
    if args.doi_list:                      # same parsing as generate_cv.main
        doi_filter = {}
        for line in open(args.doi_list, encoding='utf-8'):
            line = line.strip()
            if not line or line.startswith('#'):
                continue
            parts = [x.strip() for x in line.split(',')]
            doi_filter[parts[0].lower()] = {x.lower() for x in parts[1:] if x}
    if_override = None
    if args.if_override:
        if_override = {str(k).lower().strip(): v for k, v in
                       json.load(open(args.if_override, encoding='utf-8')).items()}

    for lg in (['en', 'ko'] if args.lang == 'both' else [args.lang]):
        generate_cv_docx(lg, doi_filter=doi_filter, pub_subtitle=args.pub_subtitle, suffix_extra=args.suffix,
                         grants_from=args.grants_from, section_breaks=not args.no_section_breaks,
                         if_override=if_override, outdir=args.outdir)


if __name__ == "__main__":
    main()
